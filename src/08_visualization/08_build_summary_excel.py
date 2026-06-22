"""
KYHPS 데이터 정제 요약 보고서(Excel) 생성
============================================
시트 구성:
1) 정제요약 (대상자수 변화·Retention)
2) 재코딩요약 (변수별 불일치/재코딩 건수)
3) Codelist_학생
4) Codelist_보호자
5) ID_5y_학생, ID_5y_보호자, ID_5y_공통
"""
import os
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

_LINUX_BASE = "/sessions/friendly-funny-cori/mnt/제1회 청소년건강패널조사"
_WIN_BASE = r"C:\Users\jiwoo\OneDrive\Desktop\나\3학년\제1회 청소년건강패널조사"
BASE = _LINUX_BASE if os.path.isdir(_LINUX_BASE) else _WIN_BASE
OUT = os.path.join(BASE, "output_cleaning")

# --- 정제 후 데이터 로드 ---
fin_stu = pd.read_csv(os.path.join(OUT, "fin_student_wide_5y.csv"), low_memory=False)
fin_par = pd.read_csv(os.path.join(OUT, "fin_parent_wide_5y.csv"), low_memory=False)
code_stu = pd.read_csv(os.path.join(OUT, "codelist_student.csv"))
code_par = pd.read_csv(os.path.join(OUT, "codelist_parent.csv"))
id_stu = pd.read_csv(os.path.join(OUT, "ID_list_student_5y.csv"))
id_par = pd.read_csv(os.path.join(OUT, "ID_list_parent_5y.csv"))
id_com = pd.read_csv(os.path.join(OUT, "ID_list_common_5y.csv"))

# --- Phase 1 요약 (실제 측정값으로) ---
WAVE_COUNTS_STU = [5051, 4702, 4438, 4409, 4243]
WAVE_COUNTS_PAR = [5051, 4689, 4424, 4388, 4182]
phase1_rows = [
    ["대상", "전체(w1)", "w2", "w3", "w4", "w5", "5년 완전응답", "Retention(%)"],
    ["학생", *WAVE_COUNTS_STU, len(id_stu), round(100*len(id_stu)/5051, 2)],
    ["보호자", *WAVE_COUNTS_PAR, len(id_par), round(100*len(id_par)/5051, 2)],
    ["공통(학생∩보호자)", "-", "-", "-", "-", "-", len(id_com), round(100*len(id_com)/5051, 2)],
]

# --- Phase 2 재코딩 요약 (파이프라인 결과 하드코딩) ---
phase2_data = pd.DataFrame([
    ["일반담배(궐련)_평생 흡연 경험(모금 기준)", 5, "1,2,3,4,5",  24,   24],
    ["일반담배(궐련)_평생 흡연 경험(개비 기준)", 5, "1,2,3,4,5",  13,   14],
    ["액상형_평생 흡연 여부(모금 기준)",          5, "1,2,3,4,5",  20,   20],
    ["궐련형_평생 흡연 경험(모금 기준)",          5, "1,2,3,4,5",  13,   16],
    ["궐련형_평생 흡연 경험(개비 기준)",          5, "1,2,3,4,5",   5,    5],
    ["평생 음주 경험(모금 기준)",                  5, "1,2,3,4,5", 1169, 2312],
    ["평생 음주 경험(잔 기준)",                    5, "1,2,3,4,5",  434,  716],
    ["평생담배제품사용경험률",                      4, "1,2,3,4",     40,   59],
], columns=["변수명(base)", "존재차시수", "차시컬럼", "불일치_subject수", "재코딩_건수(셀단위)"])

# --- Excel 작성 ---
wb = Workbook()
THIN = Side(border_style="thin", color="999999")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
HEAD_FILL = PatternFill("solid", start_color="305496")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def style_body(ws, start_row, end_row, ncols, align=CENTER):
    for r in range(start_row, end_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.alignment = align
            cell.border = BORDER


def autosize(ws, max_w=40):
    from openpyxl.utils import get_column_letter
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            try:
                col_idx = cell.column  # int
            except Exception:
                continue
            length = len(str(cell.value))
            widths[col_idx] = max(widths.get(col_idx, 10), length + 2)
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = min(w, max_w)


# Sheet 1: 정제요약
ws1 = wb.active
ws1.title = "1_정제요약"
ws1["A1"] = "KYHPS 5개년 패널 데이터 정제 요약 보고"
ws1["A1"].font = Font(name="Arial", bold=True, size=14)
ws1.merge_cells("A1:H1")
ws1["A1"].alignment = CENTER

ws1["A3"] = "[Phase 1] 차시별 응답자 수 및 5년 완전응답 Retention"
ws1["A3"].font = Font(name="Arial", bold=True, size=12)

start_r = 5
for i, row in enumerate(phase1_rows):
    for j, v in enumerate(row):
        ws1.cell(row=start_r + i, column=1 + j, value=v)
style_header(ws1, start_r, len(phase1_rows[0]))
style_body(ws1, start_r + 1, start_r + len(phase1_rows) - 1, len(phase1_rows[0]))

ws1["A11"] = "[해석 가이드]"
ws1["A11"].font = Font(name="Arial", bold=True, size=11)
notes = [
    "• 학생 Retention 78.9%, 보호자 Retention 77.7%",
    "• 5개년 모두 응답한 공통 (학생∩보호자) 페어: 3,922명",
    "• 1차 표본(5,051명) 기준 약 22% 중도탈락 → Selection Bias 검토 필요(Discussion)",
    "• 분석 단위가 학생-보호자 페어인 경우 N=3,922 확정",
]
for i, n in enumerate(notes):
    ws1.cell(row=12 + i, column=1, value=n).font = BODY_FONT
    ws1.merge_cells(start_row=12 + i, start_column=1, end_row=12 + i, end_column=8)

autosize(ws1)

# Sheet 2: 재코딩요약
ws2 = wb.create_sheet("2_재코딩요약")
ws2["A1"] = "[Phase 2] 평생 흡연/음주 변수 논리적 불일치 검사·재코딩 결과"
ws2["A1"].font = Font(name="Arial", bold=True, size=12)
ws2.merge_cells("A1:E1")

# 헤더
for j, col in enumerate(phase2_data.columns):
    ws2.cell(row=3, column=1 + j, value=col)
style_header(ws2, 3, len(phase2_data.columns))

for i, row in phase2_data.iterrows():
    for j, v in enumerate(row):
        ws2.cell(row=4 + i, column=1 + j, value=v)
style_body(ws2, 4, 4 + len(phase2_data) - 1, len(phase2_data.columns))

# 합계 행
sum_row = 4 + len(phase2_data)
ws2.cell(row=sum_row, column=1, value="합계").font = Font(name="Arial", bold=True)
ws2.cell(row=sum_row, column=4, value=f"=SUM(D4:D{sum_row-1})")
ws2.cell(row=sum_row, column=5, value=f"=SUM(E4:E{sum_row-1})")
style_body(ws2, sum_row, sum_row, 5)
ws2.cell(row=sum_row, column=1).font = Font(name="Arial", bold=True)

ws2.cell(row=sum_row + 2, column=1, value="[재코딩 로직]").font = Font(name="Arial", bold=True)
ws2.cell(row=sum_row + 3, column=1,
         value="한 번이라도 '있음(1)'이 나타나면 그 이후의 모든 차시를 '있음(1)'로 강제 재코딩 (NaN 유지)")
ws2.merge_cells(start_row=sum_row + 3, start_column=1, end_row=sum_row + 3, end_column=5)

autosize(ws2)

# Sheet 3: Codelist_학생
ws3 = wb.create_sheet("3_Codelist_학생")
for j, col in enumerate(code_stu.columns):
    ws3.cell(row=1, column=1 + j, value=col)
style_header(ws3, 1, len(code_stu.columns))
for i, row in code_stu.iterrows():
    for j, v in enumerate(row):
        ws3.cell(row=2 + i, column=1 + j, value=v if pd.notna(v) else None)
autosize(ws3, max_w=35)
ws3.freeze_panes = "A2"

# Sheet 4: Codelist_보호자
ws4 = wb.create_sheet("4_Codelist_보호자")
for j, col in enumerate(code_par.columns):
    ws4.cell(row=1, column=1 + j, value=col)
style_header(ws4, 1, len(code_par.columns))
for i, row in code_par.iterrows():
    for j, v in enumerate(row):
        ws4.cell(row=2 + i, column=1 + j, value=v if pd.notna(v) else None)
autosize(ws4, max_w=35)
ws4.freeze_panes = "A2"

# Sheet 5: ID 리스트 (학생/보호자/공통)
ws5 = wb.create_sheet("5_ID_5년완전응답")
ws5["A1"] = "학생 5년완전"
ws5["B1"] = "보호자 5년완전"
ws5["C1"] = "공통(학생∩보호자)"
style_header(ws5, 1, 3)
n_max = max(len(id_stu), len(id_par), len(id_com))
for i in range(n_max):
    if i < len(id_stu):
        ws5.cell(row=2 + i, column=1, value=id_stu.iloc[i, 0])
    if i < len(id_par):
        ws5.cell(row=2 + i, column=2, value=id_par.iloc[i, 0])
        ws5.cell(row=2 + i, column=2, value=id_par.iloc[i, 0])
    if i < len(id_com):
        ws5.cell(row=2 + i, column=3, value=id_com.iloc[i, 0])
ws5.column_dimensions["A"].width = 18
ws5.column_dimensions["B"].width = 18
ws5.column_dimensions["C"].width = 22
ws5.freeze_panes = "A2"

out_path = os.path.join(OUT, "정제요약보고_KYHPS_5y.xlsx")
wb.save(out_path)
print("saved:", out_path)
